from pptx import Presentation
import os

ppt_path = '/root/dpacwebsite/RDC_Presentation.pptx'
prs = Presentation(ppt_path)

def add_slide(title_text, bullet_points):
    try:
        bullet_slide_layout = prs.slide_layouts[1]
    except IndexError:
        bullet_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    try:
        title_shape = shapes.title
        title_shape.text = title_text
    except Exception:
        pass
    
    try:
        body_shape = None
        for shape in slide.placeholders:
            if shape.placeholder_format.idx == 1:
                body_shape = shape
                break
        if not body_shape and len(slide.placeholders) > 1:
            body_shape = slide.placeholders[1]
        
        if body_shape:
            tf = body_shape.text_frame
            tf.text = bullet_points[0]
            for bullet in bullet_points[1:]:
                p = tf.add_paragraph()
                p.text = bullet
    except Exception:
        pass

add_slide("Literature Review: Identifying the Human-Centric Gap", [
    "Zou, W., et al. (2021): Highlighted usability and domain-specific challenges in smart contract development lifecycles.",
    "Bassan, F., & Rabitti, M. (2024): Demonstrated that real-world, human-centric factors are frequently overlooked in purely technical blockchain transitions.",
    "Ahmed, S. U., et al. (2024): Proposed NLP for contract generation. Our framework extends this by introducing crucial human-in-the-loop validation mechanisms."
])

add_slide("Literature Review: Security & Architecture Foundations", [
    "Nelaturu, K., et al. (2023): Emphasized the necessity of 'correct-by-design' formal verification before deployment.",
    "Skotnica, M., et al. (2020): Advocated for a model-driven approach, justifying our use of pre-approved templates over from-scratch coding.",
    "Wei, G., et al. (2024): Supported consolidating smart contracts with behavioral logic to accurately reflect the 'intent' of legal agreements."
])

try:
    prs.save(ppt_path)
    print("Literature review slides added successfully to", ppt_path)
except Exception as e:
    print("Error updating presentation:", e)
